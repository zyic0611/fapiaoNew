#!/usr/bin/env python3
"""电子发票号码批量提取工具 — tkinter + PyMuPDF + openpyxl"""

from __future__ import annotations

import re
import threading
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import fitz
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

_VALID_INVOICE_LENGTHS = frozenset({8, 20})

_LABELED_PATTERNS = [
    re.compile(r"发票号码\s*[：:]\s*(\d{20})"),
    re.compile(r"发票号码\s*[：:]\s*(\d{8})"),
    re.compile(r"发票号码\s*[：:]\s*(\d+)"),
    re.compile(r"发票号码\s*[：:]\s*\n\s*(\d{20}|\d{8})"),
    re.compile(r"发票号码\s+(\d{20})"),
    re.compile(r"发票号码\s+(\d{8})"),
    re.compile(r"发票号码\s+(\d+)"),
]

_FALLBACK_20_DIGIT = re.compile(r"(?<![0-9])(\d{20})(?![0-9])")

MSG_NOT_RECOGNIZED = "未能识别"
MSG_FILE_ERROR = "文件异常"


def collect_pdf_paths(pdf_dir: Path) -> list[Path]:
    seen: dict[str, Path] = {}
    for pattern in ("*.pdf", "*.PDF"):
        for path in pdf_dir.glob(pattern):
            key = str(path.resolve()).lower()
            if key not in seen:
                seen[key] = path
    return sorted(seen.values(), key=lambda p: p.name.lower())


def _accept_invoice_digits(digits: str) -> str | None:
    if len(digits) in _VALID_INVOICE_LENGTHS:
        return digits
    return None


def extract_invoice_number(text: str) -> str | None:
    for pattern in _LABELED_PATTERNS:
        match = pattern.search(text)
        if match:
            accepted = _accept_invoice_digits(match.group(1))
            if accepted:
                return accepted

    match = _FALLBACK_20_DIGIT.search(text)
    if match:
        return match.group(1)
    return None


def extract_text_from_pdf(path: Path) -> str:
    with fitz.open(path) as doc:
        if doc.page_count == 0:
            return ""
        return doc[0].get_text("text")


def process_single_pdf(path: Path) -> tuple[str, str]:
    try:
        text = extract_text_from_pdf(path)
    except Exception:
        return MSG_FILE_ERROR, "skip"

    number = extract_invoice_number(text)
    if number:
        return number, "ok"
    return MSG_NOT_RECOGNIZED, "fail"


def write_excel(out_dir: Path, rows: list[tuple[str, str]]) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = out_dir / f"发票提取结果_{timestamp}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "提取结果"
    ws.append(["文件名", "发票号码"])
    for filename, number in rows:
        ws.append([filename, number])

    ws.column_dimensions[get_column_letter(1)].width = 36
    ws.column_dimensions[get_column_letter(2)].width = 24
    wb.save(out_path)
    return out_path


class InvoiceExtractorApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("电子发票号码极速提取")
        self.root.minsize(640, 480)

        self.pdf_dir_var = tk.StringVar()
        self.out_dir_var = tk.StringVar()
        self._worker: threading.Thread | None = None

        self._build_ui()

    def _build_ui(self) -> None:
        main = ttk.Frame(self.root, padding=12)
        main.pack(fill=tk.BOTH, expand=True)

        ttk.Label(main, text="PDF 文件夹：").grid(row=0, column=0, sticky=tk.W, pady=4)
        ttk.Entry(main, textvariable=self.pdf_dir_var, width=52).grid(
            row=0, column=1, sticky=tk.EW, padx=(0, 8), pady=4
        )
        ttk.Button(main, text="浏览…", command=self._browse_pdf_dir).grid(
            row=0, column=2, pady=4
        )

        ttk.Label(main, text="Excel 导出目录：").grid(row=1, column=0, sticky=tk.W, pady=4)
        ttk.Entry(main, textvariable=self.out_dir_var, width=52).grid(
            row=1, column=1, sticky=tk.EW, padx=(0, 8), pady=4
        )
        ttk.Button(main, text="浏览…", command=self._browse_out_dir).grid(
            row=1, column=2, pady=4
        )

        main.columnconfigure(1, weight=1)

        self.run_btn = ttk.Button(
            main, text="开始极速提取", command=self._start_extraction
        )
        self.run_btn.grid(row=2, column=0, columnspan=3, pady=(12, 8))

        log_frame = ttk.LabelFrame(main, text="处理日志", padding=4)
        log_frame.grid(row=3, column=0, columnspan=3, sticky=tk.NSEW, pady=(4, 0))
        main.rowconfigure(3, weight=1)
        log_frame.rowconfigure(0, weight=1)
        log_frame.columnconfigure(0, weight=1)

        self.log_text = tk.Text(
            log_frame, height=16, state=tk.DISABLED, wrap=tk.WORD, font=("", 12)
        )
        scrollbar = ttk.Scrollbar(log_frame, orient=tk.VERTICAL, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scrollbar.set)
        self.log_text.grid(row=0, column=0, sticky=tk.NSEW)
        scrollbar.grid(row=0, column=1, sticky=tk.NS)

    def _browse_pdf_dir(self) -> None:
        path = filedialog.askdirectory(title="选择 PDF 文件夹")
        if path:
            self.pdf_dir_var.set(path)

    def _browse_out_dir(self) -> None:
        path = filedialog.askdirectory(title="选择 Excel 导出目录")
        if path:
            self.out_dir_var.set(path)

    def _append_log(self, message: str) -> None:
        def _update() -> None:
            self.log_text.configure(state=tk.NORMAL)
            self.log_text.insert(tk.END, message + "\n")
            self.log_text.see(tk.END)
            self.log_text.configure(state=tk.DISABLED)

        self.root.after(0, _update)

    def _clear_log(self) -> None:
        def _update() -> None:
            self.log_text.configure(state=tk.NORMAL)
            self.log_text.delete("1.0", tk.END)
            self.log_text.configure(state=tk.DISABLED)

        self.root.after(0, _update)

    def _set_running(self, running: bool) -> None:
        def _update() -> None:
            state = tk.DISABLED if running else tk.NORMAL
            self.run_btn.configure(state=state)

        self.root.after(0, _update)

    def _validate_paths(self) -> tuple[Path, Path] | None:
        pdf_raw = self.pdf_dir_var.get().strip()
        out_raw = self.out_dir_var.get().strip()
        if not pdf_raw or not out_raw:
            messagebox.showwarning("提示", "请先选择 PDF 文件夹和 Excel 导出目录。")
            return None

        pdf_dir = Path(pdf_raw)
        out_dir = Path(out_raw)
        if not pdf_dir.is_dir():
            messagebox.showerror("错误", f"PDF 文件夹不存在：\n{pdf_dir}")
            return None
        if not out_dir.is_dir():
            messagebox.showerror("错误", f"导出目录不存在：\n{out_dir}")
            return None
        return pdf_dir, out_dir

    def _start_extraction(self) -> None:
        if self._worker and self._worker.is_alive():
            return

        paths = self._validate_paths()
        if paths is None:
            return

        pdf_dir, out_dir = paths
        self._clear_log()
        self._set_running(True)
        self._append_log(f"开始处理：{pdf_dir}")

        self._worker = threading.Thread(
            target=self._run_extraction,
            args=(pdf_dir, out_dir),
            daemon=True,
        )
        self._worker.start()

    def _run_extraction(self, pdf_dir: Path, out_dir: Path) -> None:
        success_count = 0
        rows: list[tuple[str, str]] = []

        try:
            pdf_files = collect_pdf_paths(pdf_dir)
            if not pdf_files:
                self._append_log("未找到 PDF 文件，将生成仅含表头的 Excel。")
            else:
                self._append_log(f"共发现 {len(pdf_files)} 个 PDF 文件。")

            for path in pdf_files:
                number, status = process_single_pdf(path)
                rows.append((path.name, number))

                if status == "skip":
                    self._append_log(f"跳过：文件异常 — {path.name}")
                elif status == "ok":
                    success_count += 1
                    self._append_log(f"{path.name} -> {number}")
                else:
                    self._append_log(f"{path.name} -> {number}")

            out_path = write_excel(out_dir, rows)
            self._append_log(f"完成。成功识别 {success_count} / {len(rows)} 个。")
            self._append_log(f"已导出：{out_path}")

            def _notify() -> None:
                messagebox.showinfo(
                    "提取完成",
                    f"共处理 {len(rows)} 个文件，成功识别 {success_count} 个。\n\n"
                    f"结果已保存至：\n{out_path}",
                )

            self.root.after(0, _notify)
        except Exception as exc:
            self._append_log(f"导出失败：{exc}")

            def _err() -> None:
                messagebox.showerror("错误", f"处理过程中发生错误：\n{exc}")

            self.root.after(0, _err)
        finally:
            self._set_running(False)


def main() -> None:
    root = tk.Tk()
    InvoiceExtractorApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
